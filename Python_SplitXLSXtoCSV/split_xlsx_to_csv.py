#!/usr/bin/env python3
"""
Split XLSX to CSV — Multi-Sheet Edition
========================================
Baca file Excel (.xlsx), tiap sheet jadi CSV terpisah.
Pilih sheet mana yang di-export, pilih kolom per sheet.

Cara pakai: python split_xlsx_to_csv.py

Author: Hermes Agent for Rholly
"""

import os
import sys
import csv
import zipfile

# ─── Helper: warna CLI ───
class c:
    HEADER = '\033[95m'
    OK = '\033[92m'
    WARN = '\033[93m'
    ERR = '\033[91m'
    CYAN = '\033[96m'
    BOLD = '\033[1m'
    END = '\033[0m'


def banner():
    print(f"""
{c.CYAN}{c.BOLD}╔══════════════════════════════════════════════╗
║   Split XLSX to CSV — Multi-Sheet Edition    ║
╚══════════════════════════════════════════════╝{c.END}
""")


def ask_file_path() -> str:
    while True:
        path = input(f"{c.BOLD}📂 Masukkan path file Excel (.xlsx/.xls): {c.END}").strip()
        if not path:
            print(f"{c.WARN}  ⚠ Path tidak boleh kosong.{c.END}")
            continue
        if not os.path.isfile(path):
            print(f"{c.ERR}  ❌ File tidak ditemukan: {path}{c.END}")
            continue
        if not path.lower().endswith(('.xlsx', '.xlsm', '.xls')):
            print(f"{c.WARN}  ⚠ File bukan .xlsx/.xls. Yakin lanjut? (y/n): {c.END}", end=' ')
            if input().strip().lower() != 'y':
                continue
        return os.path.abspath(path)


def detect_format(filepath: str) -> str:
    if zipfile.is_zipfile(filepath):
        return 'xlsx'
    try:
        with open(filepath, 'rb') as f:
            magic = f.read(8)
        if magic == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
            return 'xls'
    except Exception:
        pass
    return 'unknown'


def load_workbook_all(filepath: str):
    """Load workbook, return (wb, format, sheet_names)."""
    fmt = detect_format(filepath)

    if fmt == 'xlsx':
        try:
            from openpyxl import load_workbook
        except ImportError:
            print(f"{c.ERR}  ❌ Package 'openpyxl' belum terpasang. Install: pip install openpyxl{c.END}")
            sys.exit(1)
        wb = load_workbook(filepath, read_only=True, data_only=True)
        return wb, 'xlsx', wb.sheetnames

    elif fmt == 'xls':
        try:
            import xlrd
        except ImportError:
            print(f"{c.ERR}  ❌ Package 'xlrd' belum terpasang. Install: pip install xlrd{c.END}")
            sys.exit(1)
        wb = xlrd.open_workbook(filepath)
        return wb, 'xls', wb.sheet_names()

    else:
        print(f"{c.ERR}  ❌ Format file tidak dikenali.{c.END}")
        sys.exit(1)


def get_sheet_info(wb, fmt, sheet_names):
    """Return list of (name, rows, cols) per sheet."""
    info = []
    for name in sheet_names:
        if fmt == 'xlsx':
            ws = wb[name]
            rows = ws.max_row or 0
            cols = ws.max_column or 0
        else:  # xls
            ws = wb.sheet_by_name(name)
            rows = ws.nrows
            cols = ws.ncols
        info.append((name, rows, cols))
    return info


def select_sheets(sheet_info):
    """Tampilkan daftar sheet, user pilih mana yang di-export."""
    print(f"\n{c.BOLD}📋 Daftar Sheet:{c.END}")
    for i, (name, rows, cols) in enumerate(sheet_info):
        print(f"  {c.CYAN}[{i}]{c.END} {name}  ({rows:,} baris × {cols} kolom)")

    print(f"\n  {c.CYAN}all{c.END} = semua sheet")
    print(f"  {c.CYAN}0,2,4{c.END} = sheet tertentu (comma-separated)")
    print(f"  {c.CYAN}0-3{c.END} = range sheet")

    while True:
        choice = input(f"\n{c.BOLD}Pilih sheet (default: all): {c.END}").strip()
        if not choice or choice.lower() == 'all':
            return list(range(len(sheet_info)))

        indices = set()
        for part in choice.split(','):
            part = part.strip()
            if '-' in part:
                try:
                    a, b = part.split('-')
                    for n in range(int(a), int(b) + 1):
                        if 0 <= n < len(sheet_info):
                            indices.add(n)
                except ValueError:
                    pass
            else:
                try:
                    n = int(part)
                    if 0 <= n < len(sheet_info):
                        indices.add(n)
                except ValueError:
                    pass

        if indices:
            return sorted(indices)
        print(f"{c.WARN}  ⚠ Input tidak valid.{c.END}")


def get_headers(wb, fmt, sheet_name):
    """Baca header row dari sheet."""
    if fmt == 'xlsx':
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            return [str(cell) if cell is not None else f"Col_{i+1}" for i, cell in enumerate(row)]
        return []
    else:
        ws = wb.sheet_by_name(sheet_name)
        if ws.nrows == 0:
            return []
        return [str(ws.cell_value(0, i)) if ws.cell_value(0, i) != '' else f"Col_{i+1}"
                for i in range(ws.ncols)]


def show_columns(headers, sheet_name):
    """Tampilkan kolom, user pilih mana yang di-export."""
    print(f"\n{c.BOLD}📋 Sheet: {sheet_name} — {len(headers)} kolom:{c.END}")
    for i, h in enumerate(headers):
        print(f"  {c.CYAN}[{i}]{c.END} {h}")

    print(f"\n  {c.CYAN}all{c.END} = semua kolom")
    print(f"  {c.CYAN}0,2,4{c.END} = kolom tertentu")
    print(f"  {c.CYAN}0-3{c.END} = range kolom")
    print(f"  {c.CYAN}skip{c.END} = lewati sheet ini")

    while True:
        choice = input(f"{c.BOLD}Pilih kolom (default: all): {c.END}").strip()
        if choice.lower() == 'skip':
            return None
        if not choice or choice.lower() == 'all':
            return list(range(len(headers)))

        indices = set()
        for part in choice.split(','):
            part = part.strip()
            if '-' in part:
                try:
                    a, b = part.split('-')
                    for n in range(int(a), int(b) + 1):
                        if 0 <= n < len(headers):
                            indices.add(n)
                except ValueError:
                    pass
            else:
                try:
                    n = int(part)
                    if 0 <= n < len(headers):
                        indices.add(n)
                except ValueError:
                    pass

        if indices:
            return sorted(indices)
        print(f"{c.WARN}  ⚠ Input tidak valid.{c.END}")


def ask_split_size(total_rows: int) -> int:
    """Tanya user mau split per berapa baris. 0 = tidak split."""
    print(f"\n{c.BOLD}📊 Total baris: {total_rows:,}{c.END}")
    print(f"  {c.CYAN}0{c.END} = export semua ke 1 file CSV")
    print(f"  {c.CYAN}10000{c.END} = split per 10.000 baris")

    while True:
        val = input(f"{c.BOLD}Split per berapa baris? (default: 0): {c.END}").strip()
        if not val:
            return 0
        try:
            n = int(val)
            if n >= 0:
                return n
        except ValueError:
            pass
        print(f"{c.WARN}  ⚠ Input tidak valid.{c.END}")


def ask_output_dir(filepath: str) -> str:
    """Default output dir = same folder as input file."""
    default = os.path.dirname(filepath)
    val = input(f"{c.BOLD}📂 Output folder (default: {default}): {c.END}").strip()
    out = val if val else default
    os.makedirs(out, exist_ok=True)
    return out


def sanitize_filename(name: str) -> str:
    """Bersihkan nama untuk filesystem."""
    return name.replace('/', '_').replace('\\', '_').replace(' ', '_').replace(':', '')


def export_sheet_to_csv(wb, fmt, sheet_name, headers, col_indices,
                        split_size, out_dir, base_name):
    """Export satu sheet ke CSV file(s)."""
    selected = [headers[i] for i in col_indices]

    if split_size == 0:
        # Single file
        fname = f"{base_name}_{sanitize_filename(sheet_name)}.csv"
        fpath = os.path.join(out_dir, fname)
        row_count = _write_single(wb, fmt, sheet_name, col_indices, selected, fpath)
        print(f"  {c.OK}✓ {fname} ({row_count:,} baris){c.END}")
        return [(fname, row_count)]

    # Split per N rows
    files = []
    part = 0
    row_count = 0
    current_rows = []

    for row in _iter_data_rows(wb, fmt, sheet_name, col_indices):
        current_rows.append(row)
        row_count += 1
        if len(current_rows) >= split_size:
            fname = f"{base_name}_{sanitize_filename(sheet_name)}_part{part+1}.csv"
            fpath = os.path.join(out_dir, fname)
            _write_csv(fpath, selected, current_rows)
            print(f"  {c.OK}✓ {fname} ({len(current_rows):,} baris){c.END}")
            files.append((fname, len(current_rows)))
            current_rows = []
            part += 1

    # Sisa
    if current_rows:
        fname = f"{base_name}_{sanitize_filename(sheet_name)}_part{part+1}.csv"
        fpath = os.path.join(out_dir, fname)
        _write_csv(fpath, selected, current_rows)
        print(f"  {c.OK}✓ {fname} ({len(current_rows):,} baris){c.END}")
        files.append((fname, len(current_rows)))

    return files


def _iter_data_rows(wb, fmt, sheet_name, col_indices):
    """Yield data rows (skip header) as list of selected column values."""
    if fmt == 'xlsx':
        ws = wb[sheet_name]
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
            yield [row[idx] if idx < len(row) else '' for idx in col_indices]
    else:
        ws = wb.sheet_by_name(sheet_name)
        for r in range(1, ws.nrows):
            yield [ws.cell_value(r, idx) for idx in col_indices]


def _write_single(wb, fmt, sheet_name, col_indices, headers, fpath):
    """Write all rows to a single CSV."""
    with open(fpath, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        count = 0
        for row in _iter_data_rows(wb, fmt, sheet_name, col_indices):
            writer.writerow(row)
            count += 1
    return count


def _write_csv(fpath, headers, rows):
    """Write rows to CSV file."""
    with open(fpath, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)


def main():
    banner()

    # 1. File path
    filepath = ask_file_path()
    base_name = os.path.splitext(os.path.basename(filepath))[0]

    # 2. Load workbook + detect format
    wb, fmt, sheet_names = load_workbook_all(filepath)
    print(f"\n{c.OK}  ✓ Format: {fmt.upper()} | {len(sheet_names)} sheet{c.END}")

    # 3. Get sheet info (rows × cols)
    sheet_info = get_sheet_info(wb, fmt, sheet_names)

    # 4. Select sheets
    selected_sheets = select_sheets(sheet_info)

    # 5. Output dir
    out_dir = ask_output_dir(filepath)

    # 6. Per-sheet: select columns + split size + export
    print(f"\n{c.BOLD}{'═'*50}{c.END}")
    print(f"{c.BOLD}  Memproses {len(selected_sheets)} sheet...{c.END}")
    print(f"{c.BOLD}{'═'*50}{c.END}")

    all_results = []

    for idx in selected_sheets:
        sheet_name = sheet_info[idx][0]
        total_rows = sheet_info[idx][1]
        print(f"\n{c.CYAN}{c.BOLD}── Sheet [{idx}]: {sheet_name} ({total_rows:,} baris) ──{c.END}")

        # Headers
        headers = get_headers(wb, fmt, sheet_name)
        if not headers:
            print(f"  {c.WARN}⚠ Sheet kosong, skip.{c.END}")
            continue

        # Select columns
        col_indices = show_columns(headers, sheet_name)
        if col_indices is None:
            print(f"  {c.WARN}⏭ Sheet dilewati.{c.END}")
            continue

        # Split size
        split_size = ask_split_size(total_rows)

        # Export
        results = export_sheet_to_csv(wb, fmt, sheet_name, headers, col_indices,
                                       split_size, out_dir, base_name)
        all_results.extend(results)

    # 7. Summary
    print(f"\n{c.BOLD}{'═'*50}{c.END}")
    print(f"{c.OK}{c.BOLD}  Selesai! {len(all_results)} file CSV dibuat:{c.END}")
    print(f"{c.BOLD}{'═'*50}{c.END}")
    for fname, rows in all_results:
        print(f"  📄 {fname}  ({rows:,} baris)")
    print(f"\n{c.CYAN}  Folder: {out_dir}{c.END}\n")


if __name__ == '__main__':
    try:
        main()
    except KeyboardInterrupt:
        print(f"\n{c.WARN}  ⏹ Dibatalkan.{c.END}")
        sys.exit(0)
